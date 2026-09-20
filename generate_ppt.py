import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_pitch_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme colors
    BG_COLOR = RGBColor(7, 11, 20)       # #070b14 Dark Navy
    CARD_BG = RGBColor(15, 23, 42)       # #0f172a Slate 900
    ACCENT_BLUE = RGBColor(59, 130, 246) # #3b82f6 Blue
    ACCENT_PURPLE = RGBColor(168, 85, 247) # #a855f7 Purple
    ACCENT_GREEN = RGBColor(16, 185, 129) # #10b981 Emerald
    TEXT_WHITE = RGBColor(248, 250, 252) # #f8fafc
    TEXT_MUTED = RGBColor(148, 163, 184) # #94a3b8 Slate 400
    BORDER_COLOR = RGBColor(30, 41, 59)  # Slate 800

    def set_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_tag):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(0.35))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_tag.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_BLUE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.7))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_background(slide1)

    # Main Title Container
    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(3.5))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "HACK DEVENGERS 2.0 • OPEN INNOVATION TRACK"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = ACCENT_BLUE
    p_badge.space_after = Pt(14)

    p1 = tf1.add_paragraph()
    p1.text = "LLM Reliability Lab"
    p1.font.size = Pt(50)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "Automated AI Model Stress-Testing, Hallucination Detection & Security Red-Teaming Arena"
    p2.font.size = Pt(20)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(28)

    # Footer Info
    foot_box = slide1.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(11.333), Inches(1.0))
    tf_foot = foot_box.text_frame
    p_foot1 = tf_foot.paragraphs[0]
    p_foot1.text = "Built by: Prateek Yadav (TechPrateek) • Hack Devengers 2.0"
    p_foot1.font.size = Pt(13)
    p_foot1.font.color.rgb = TEXT_WHITE
    p_foot1.font.bold = True

    p_foot2 = tf_foot.add_paragraph()
    p_foot2.text = "Live Demo: https://llm-reliability-lab-seven.vercel.app  |  GitHub: https://github.com/TechPrateek/LLM-Reliability-Lab"
    p_foot2.font.size = Pt(12)
    p_foot2.font.color.rgb = ACCENT_BLUE

    # -------------------------------------------------------------------------
    # SLIDE 2: The Problem
    # -------------------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_background(slide2)
    add_header(slide2, "The Problem: Why Blind LLM Deployment Fails in Production", "01. PROBLEM STATEMENT")

    cards_data2 = [
        ("🚨 Hallucination Traps", "Models fall for deceptive or counterfactual premises (e.g. 'Da Vinci's iPhone in 1503') and generate confident, fabricated claims that erode user trust."),
        ("🔓 Jailbreak & Delimiter Injection", "Users exploit persona prompts (like DAN v12) or delimiter injections ([ADMIN OVERRIDE]) to bypass guardrails and leak internal system prompts."),
        ("💥 Schema & Constraint Drift", "Production microservices need clean JSON. Models frequently include conversational markdown fences (```json) or violate negative constraints, breaking downstream pipelines."),
        ("⏱️ Hidden Latency & Token Economics", "Engineers lack unified side-by-side visibility into Time to First Token (TTFT), real token throughput, and exact query cost across competing LLM architectures.")
    ]

    for idx, (title, desc) in enumerate(cards_data2):
        col = idx % 2
        row = idx // 2
        left = Inches(1.0 + col * 5.8)
        top = Inches(1.8 + row * 2.5)

        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.4), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tb = slide2.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(5.0), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------------------
    # SLIDE 3: The Solution
    # -------------------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_background(slide3)
    add_header(slide3, "The Solution: LLM Reliability Lab Arena & Test CI", "02. ARCHITECTURE & SOLUTION")

    cards_data3 = [
        ("⚔️ Multi-Model Battle Arena", "Compare 2 to 3 frontier models (Gemini 2.5 Flash, GPT-4o Mini, Llama 3.3, Claude 3.5, DeepSeek R1) simultaneously with client-side TTFT, latency, and query cost profiling."),
        ("🧪 Automated Stress-Testing Battery", "Pre-engineered test vectors across 4 failure categories: Hallucinations, Red-Teaming Jailbreaks, Strict Schema Formatting, and Counterfactual Reasoning."),
        ("📊 Model Reliability Index (MRI™)", "Algorithmic grading engine (0–100) that evaluates truthfulness, security defense, formatting compliance, and latency performance."),
        ("🏷️ AI Safety Nutrition Labels", "One-click exportable official audit sheets (PDF & Markdown) formatted like FDA nutrition facts, mapped to NIST AI RMF and EU AI Act baselines.")
    ]

    for idx, (title, desc) in enumerate(cards_data3):
        col = idx % 2
        row = idx // 2
        left = Inches(1.0 + col * 5.8)
        top = Inches(1.8 + row * 2.5)

        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.4), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_BLUE

        tb = slide3.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(5.0), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_BLUE
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------------------
    # SLIDE 4: Uniqueness & Competitive Edge
    # -------------------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_background(slide4)
    add_header(slide4, "Uniqueness & Competitive Edge: Why We Stand Out", "03. INNOVATION & VALUE")

    cards_data4 = [
        ("1. FDA-Style Nutrition Facts for AI", "While other tools spit out confusing log files, we generate standardized, executive-ready 1-page PDF compliance labels showing risk percentages and operational economics."),
        ("2. Multi-Axis Geometric Spider Graphs", "5-axis radar chart showing normalized trade-offs across Truthfulness, Jailbreak Defense, Schema Adherence, Speed, and Cost."),
        ("3. Zero-Barrier Judging Engine", "Hackathon judges can test all features, diff views, and export cards immediately without needing an API key, or connect live Gemini keys seamlessly."),
        ("4. Strategy Diffing & CoT Trace", "Direct textual diff between competing models and dedicated collapsible view for chain-of-thought reasoning models (e.g. DeepSeek R1).")
    ]

    for idx, (title, desc) in enumerate(cards_data4):
        col = idx % 2
        row = idx // 2
        left = Inches(1.0 + col * 5.8)
        top = Inches(1.8 + row * 2.5)

        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.4), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_PURPLE

        tb = slide4.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(5.0), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_PURPLE
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------------------
    # SLIDE 5: Tech Stack & Architecture
    # -------------------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_background(slide5)
    add_header(slide5, "Technology Stack & Production Engineering", "04. TECH SPECS")

    stack_items = [
        ("Frontend & App Framework", "React 19, TypeScript, Vite, Tailwind CSS v4, Lucide Icons, Framer Motion"),
        ("Data Visualization Engine", "Recharts (Multi-axis Radar Charts, Latency Bar Charts, Pareto Matrices)"),
        ("Client-Side PDF & Audit Engine", "jsPDF & html2canvas (In-browser canvas vector rendering for instant PDF export)"),
        ("Frontier AI Models Evaluated", "Google Gemini 2.5 Flash, OpenAI GPT-4o Mini, Meta Llama 3.3 70B, Claude 3.5, DeepSeek R1"),
        ("Deployment & DevOps", "Vercel Edge Network, Git, GitHub Actions, 100% Client-Side Private Storage")
    ]

    for idx, (label, details) in enumerate(stack_items):
        top = Inches(1.7 + idx * 1.05)
        bar = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), top, Inches(11.333), Inches(0.9))
        bar.fill.solid()
        bar.fill.fore_color.rgb = CARD_BG
        bar.line.color.rgb = BORDER_COLOR

        tb = slide5.shapes.add_textbox(Inches(1.3), top + Inches(0.12), Inches(10.7), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True

        p_l = tf.paragraphs[0]
        p_l.text = label
        p_l.font.size = Pt(13)
        p_l.font.bold = True
        p_l.font.color.rgb = ACCENT_GREEN

        p_det = tf.add_paragraph()
        p_det.text = details
        p_det.font.size = Pt(11)
        p_det.font.color.rgb = TEXT_WHITE

    # -------------------------------------------------------------------------
    # SLIDE 6: Summary & Live Links
    # -------------------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_background(slide6)
    add_header(slide6, "Try LLM Reliability Lab Live Today", "05. CONCLUSION & DEMO")

    # Big Central Card
    card_sum = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(11.333), Inches(5.0))
    card_sum.fill.solid()
    card_sum.fill.fore_color.rgb = CARD_BG
    card_sum.line.color.rgb = ACCENT_BLUE

    tb_s = slide6.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(4.2))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True

    p_s1 = tf_s.paragraphs[0]
    p_s1.text = "🎯 Hackathon Deliverables & Quick Links"
    p_s1.font.size = Pt(22)
    p_s1.font.bold = True
    p_s1.font.color.rgb = TEXT_WHITE
    p_s1.space_after = Pt(14)

    items6 = [
        ("🌐 Live Deployment URL:", "https://llm-reliability-lab-seven.vercel.app"),
        ("💻 GitHub Repository:", "https://github.com/TechPrateek/LLM-Reliability-Lab"),
        ("⚡ Key Highlights:", "Full stress-testing battery, side-by-side arena, 5-axis spider radar, and downloadable 1-page PDF Nutrition Label."),
        ("👥 Target Audience:", "AI Engineers, Red-Team Security Researchers, and Enterprise Compliance Teams.")
    ]

    for label, val in items6:
        p_item = tf_s.add_paragraph()
        p_item.text = f"{label} {val}"
        p_item.font.size = Pt(14)
        p_item.font.color.rgb = TEXT_WHITE
        p_item.space_after = Pt(10)

    # Save presentation
    output_path = "LLM_Reliability_Lab_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_pitch_deck()
